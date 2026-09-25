"""
Script de ETL (Extract, Transform, Load) para Sistema de Gestão de Frotas (SGFrotas)
Lê as planilhas reais em docs/Planilhas/ e gera a migração V7__seed_initial_data.sql para o Flyway.
"""

import os
import re
import csv
from datetime import datetime
import openpyxl

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS_PLANILHAS = os.path.join(BASE_DIR, "docs", "Planilhas")
OUTPUT_SQL = os.path.join(BASE_DIR, "db", "migrations", "V7__seed_initial_data.sql")

def escape_sql(val):
    if val is None:
        return "NULL"
    val_str = str(val).strip()
    if val_str == "" or val_str.lower() == "none" or val_str.lower() == "null":
        return "NULL"
    cleaned = val_str.replace("'", "''")
    return f"'{cleaned}'"

def format_date_sql(val):
    if not val:
        return "NULL"
    if isinstance(val, datetime):
        return f"'{val.strftime('%Y-%m-%d')}'"
    val_str = str(val).strip()
    # Tenta padrão DD/MM/YYYY
    match_br = re.match(r'^(\d{2})/(\d{2})/(\d{4})', val_str)
    if match_br:
        d, m, y = match_br.groups()
        return f"'{y}-{m}-{d}'"
    # Tenta padrão YYYY-MM-DD
    match_iso = re.match(r'^(\d{4})-(\d{2})-(\d{2})', val_str)
    if match_iso:
        return f"'{val_str[:10]}'"
    return "NULL"

def format_timestamp_sql(val):
    if not val:
        return "NULL"
    if isinstance(val, datetime):
        return f"'{val.strftime('%Y-%m-%d %H:%M:%S+00')}'"
    val_str = str(val).strip()
    match_br = re.match(r'^(\d{2})/(\d{2})/(\d{4})\s*(\d{2}:\d{2}(?::\d{2})?)?', val_str)
    if match_br:
        d, m, y, hora = match_br.groups()
        h = hora if hora else "00:00:00"
        if len(h) == 5:
            h += ":00"
        return f"'{y}-{m}-{d} {h}+00'"
    match_iso = re.match(r'^(\d{4})-(\d{2})-(\d{2})\s*(\d{2}:\d{2}(?::\d{2})?)?', val_str)
    if match_iso:
        d_str = val_str[:10]
        h_part = val_str[11:].strip() if len(val_str) > 10 else "00:00:00"
        return f"'{d_str} {h_part}+00'"
    return "NULL"

def clean_decimal(val, default=0.0):
    if val is None:
        return str(default)
    val_str = str(val).strip()
    if not val_str:
        return str(default)
    # Trata formato brasileiro 1.933,53 ou americano 1933.53
    val_str = val_str.replace("R$", "").replace(" ", "")
    if "," in val_str and "." in val_str:
        val_str = val_str.replace(".", "").replace(",", ".")
    elif "," in val_str:
        val_str = val_str.replace(",", ".")
    try:
        f = float(val_str)
        return f"{f:.2f}"
    except ValueError:
        return str(default)

def clean_cpf(val):
    if not val:
        return None
    val_str = re.sub(r'\D', '', str(val))
    if len(val_str) == 11:
        return f"{val_str[:3]}.{val_str[3:6]}.{val_str[6:9]}-{val_str[9:]}"
    return str(val).strip()[:14] if str(val).strip() else None

def clean_placa(val):
    if not val:
        return None
    placa = str(val).strip().upper().replace("-", "")
    return placa if len(placa) >= 7 else None

def run_etl():
    print("Iniciando processo ETL para SGFrotas...", flush=True)
    sql_statements = [
        "-- =====================================================================",
        "-- Migration V7: Seed de Dados Iniciais Migrados das Planilhas Legadas",
        "-- Database: PostgreSQL 16 (Flyway Migrations)",
        "-- =====================================================================\n"
    ]

    planilha_controle = os.path.join(DOCS_PLANILHAS, "Planilha Controle Frota.xlsx")
    print(f"Carregando {planilha_controle} em modo read_only...", flush=True)
    wb_controle = openpyxl.load_workbook(planilha_controle, read_only=True, data_only=True)
    print("Planilha de controle carregada!", flush=True)

    # 1. Locadora Padrão
    sql_statements.append("-- 1. Locadoras")
    sql_statements.append("INSERT INTO locadoras (nome, contato_suporte) VALUES ('ARVAL', 'suporte@arval.com.br') ON CONFLICT (nome) DO NOTHING;\n")

    # 2. Centros de Custo
    sql_statements.append("-- 2. Centros de Custo")
    centros_custo_map = set()
    if 'CC' in wb_controle.sheetnames:
        ws_cc = wb_controle['CC']
        for row in ws_cc.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                cod = str(row[0]).strip()
                op = str(row[1]).strip() if len(row) > 1 and row[1] else f"Operação {cod}"
                if cod not in centros_custo_map:
                    centros_custo_map.add(cod)
                    sql_statements.append(f"INSERT INTO centros_custo (codigo, nome_operacao) VALUES ({escape_sql(cod)}, {escape_sql(op)}) ON CONFLICT (codigo) DO NOTHING;")
    sql_statements.append("")

    # 3. ATPs e Gestores
    sql_statements.append("-- 3. Assistências Técnicas Positivo (ATPs)")
    atps_map = set()
    if 'Gestores' in wb_controle.sheetnames:
        ws_gestores = wb_controle['Gestores']
        for row in ws_gestores.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                cod_atp = str(row[0]).strip()
                op = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                nome_atp = str(row[2]).strip() if len(row) > 2 and row[2] else f"ATP {cod_atp}"
                uf = str(row[3]).strip()[:2].upper() if len(row) > 3 and row[3] else "PR"
                sup = str(row[4]).strip() if len(row) > 4 and row[4] else None
                coord = str(row[5]).strip() if len(row) > 5 and row[5] else None
                
                if cod_atp not in atps_map:
                    atps_map.add(cod_atp)
                    sql_statements.append(
                        f"INSERT INTO atps (codigo_atp, nome, uf, supervisor_nome, coordenador_nome) "
                        f"VALUES ({escape_sql(cod_atp)}, {escape_sql(nome_atp)}, {escape_sql(uf)}, {escape_sql(sup)}, {escape_sql(coord)}) "
                        f"ON CONFLICT (codigo_atp) DO UPDATE SET nome = EXCLUDED.nome, supervisor_nome = EXCLUDED.supervisor_nome;"
                    )
    sql_statements.append("")

    # 4. Técnicos & Condutores
    sql_statements.append("-- 4. Técnicos de Campo e Condutores")
    tecnicos_map = set()
    if 'Base de técnicos' in wb_controle.sheetnames:
        ws_tecnicos = wb_controle['Base de técnicos']
        # Cabeçalho está na linha 2 (index 1), dados a partir da linha 3 (min_row=3)
        for row in ws_tecnicos.iter_rows(min_row=3, values_only=True):
            if not row or not row[1]:
                continue
            matr = str(row[1]).strip()
            nome = str(row[2]).strip() if len(row) > 2 and row[2] else "NÃO INFORMADO"
            status = str(row[3]).strip() if len(row) > 3 and row[3] else "Ativo"
            cargo = str(row[4]).strip() if len(row) > 4 and row[4] else None
            cc_cod = str(row[5]).strip() if len(row) > 5 and row[5] else None
            admissao = format_date_sql(row[6]) if len(row) > 6 else "NULL"
            rg = str(row[12]).strip() if len(row) > 12 and row[12] else None
            cpf = clean_cpf(row[13]) if len(row) > 13 else None
            nascimento = format_date_sql(row[14]) if len(row) > 14 else "NULL"
            sap = str(row[15]).strip() if len(row) > 15 and row[15] else None

            if matr and matr not in tecnicos_map:
                tecnicos_map.add(matr)
                # Garante centro de custo se referenciado
                if cc_cod and cc_cod not in centros_custo_map:
                    centros_custo_map.add(cc_cod)
                    sql_statements.append(f"INSERT INTO centros_custo (codigo, nome_operacao) VALUES ({escape_sql(cc_cod)}, 'CC {cc_cod}') ON CONFLICT (codigo) DO NOTHING;")

                cc_subquery = f"(SELECT id FROM centros_custo WHERE codigo = {escape_sql(cc_cod)} LIMIT 1)" if cc_cod else "NULL"
                sql_statements.append(
                    f"INSERT INTO tecnicos (matricula, nome, status_colaborador, cargo, centro_custo_id, data_admissao, rg, cpf, data_nascimento, codigo_sap_fornecedor) "
                    f"VALUES ({escape_sql(matr)}, {escape_sql(nome)}, {escape_sql(status)}, {escape_sql(cargo)}, {cc_subquery}, {admissao}, {escape_sql(rg)}, {escape_sql(cpf)}, {nascimento}, {escape_sql(sap)}) "
                    f"ON CONFLICT (matricula) DO NOTHING;"
                )
    sql_statements.append("")

    # 5. Veículos
    sql_statements.append("-- 5. Veículos da Frota")
    veiculos_map = {}
    if 'Base Veiculos x técnicos' in wb_controle.sheetnames:
        ws_veic = wb_controle['Base Veiculos x técnicos']
        for row in ws_veic.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            placa = clean_placa(row[1])
            if not placa:
                continue
            
            data_entrega = format_date_sql(row[0]) if row[0] else "NULL"
            tecnico_nome = str(row[2]).strip() if len(row) > 2 and row[2] else None
            filial = str(row[4]).strip() if len(row) > 4 and row[4] else ""
            base = str(row[5]).strip() if len(row) > 5 and row[5] else ""
            hodo = clean_decimal(row[7], default=0.0) if len(row) > 7 else "0.00"
            franquia_mes = clean_decimal(row[12], default=2800.0) if len(row) > 12 else "2800.00"
            franquia_fim = clean_decimal(row[13], default=100800.0) if len(row) > 13 else "100800.00"

            # Tenta associar com ATP existente
            atp_subquery = f"(SELECT id FROM atps WHERE nome ILIKE '%{base}%' OR nome ILIKE '%{filial}%' LIMIT 1)" if (base or filial) else "NULL"
            locadora_subquery = "(SELECT id FROM locadoras WHERE nome = 'ARVAL' LIMIT 1)"

            veiculos_map[placa] = {
                "tecnico_nome": tecnico_nome,
                "hodometro": hodo,
                "data_entrega": data_entrega
            }

            sql_statements.append(
                f"INSERT INTO veiculos (placa, modelo, locadora_id, atp_id, status_operacional, hodometro_atual, franquia_mensal_km, franquia_final_km, data_inicio_contrato) "
                f"VALUES ({escape_sql(placa)}, 'RENAULT KWID ZEN 1.0', {locadora_subquery}, {atp_subquery}, 'EM_USO', {hodo}, {franquia_mes}, {franquia_fim}, {data_entrega}) "
                f"ON CONFLICT (placa) DO UPDATE SET hodometro_atual = EXCLUDED.hodometro_atual, updated_at = clock_timestamp();"
            )
    sql_statements.append("")

    # 6. Custódias de Veículos (Vínculo Veículo x Condutor)
    sql_statements.append("-- 6. Custódia / Atribuição de Condutores aos Veículos")
    for placa, info in veiculos_map.items():
        if info["tecnico_nome"] and info["tecnico_nome"] != "0":
            tec_clean = info["tecnico_nome"].replace("'", "''")
            data_ini = info["data_entrega"] if info["data_entrega"] != "NULL" else "CURRENT_DATE"
            hodo = info["hodometro"]
            sql_statements.append(
                f"INSERT INTO custodias_veiculos (veiculo_id, tecnico_id, data_retirada, hodometro_retirada, status) "
                f"SELECT v.id, t.id, COALESCE({data_ini}, CURRENT_DATE), {hodo}, 'ATIVA' "
                f"FROM veiculos v, tecnicos t "
                f"WHERE v.placa = '{placa}' AND t.nome ILIKE '%{tec_clean}%' "
                f"LIMIT 1;"
            )
    sql_statements.append("")

    # 7. Manutenções Preventivas
    sql_statements.append("-- 7. Manutenções Preventivas e Revisões")
    planilha_manutencao = os.path.join(DOCS_PLANILHAS, "manutenção preventiva .xlsx")
    if os.path.exists(planilha_manutencao):
        print(f"Carregando {planilha_manutencao} em modo read_only...", flush=True)
        wb_manut = openpyxl.load_workbook(planilha_manutencao, read_only=True, data_only=True)
        if 'condutores' in wb_manut.sheetnames:
            ws_man = wb_manut['condutores']
            for row in ws_man.iter_rows(min_row=2, values_only=True):
                if not row or not row[1]:
                    continue
                placa = clean_placa(row[1])
                if not placa:
                    continue
                km_real = clean_decimal(row[3], default=0.0)
                km_ultima = clean_decimal(row[4], default=0.0)
                km_prox = clean_decimal(row[5], default=0.0)
                status_raw = str(row[7]).strip().upper() if len(row) > 7 and row[7] else "OK"
                data_agendada = format_date_sql(row[8]) if len(row) > 8 else "NULL"
                obs = str(row[9]).strip() if len(row) > 9 and row[9] else ""

                status_map = {
                    "OK": "CONCLUIDA",
                    "URGENTE": "PENDENTE",
                    "AGENDADA": "AGENDADA",
                    "PENDENTE": "PENDENTE"
                }
                status_man = status_map.get(status_raw, "PENDENTE")
                ciclo_km = int(float(km_prox)) if float(km_prox) > 0 else None
                ciclo_val = str(ciclo_km) if ciclo_km else "NULL"

                sql_statements.append(
                    f"INSERT INTO manutencoes_veiculos (veiculo_id, tipo_manutencao, ciclo_km, odometro_momento, proxima_revisao_km, data_agendamento, status, notas) "
                    f"SELECT v.id, 'PREVENTIVA', {ciclo_val}, {km_real}, {km_prox}, {data_agendada}, '{status_man}', {escape_sql(obs)} "
                    f"FROM veiculos v WHERE v.placa = '{placa}' LIMIT 1;"
                )
    sql_statements.append("")

    # 8. Infrações e Multas Históricas
    sql_statements.append("-- 8. Histórico de Infrações e Multas")
    if 'Infrações ' in wb_controle.sheetnames:
        ws_infr = wb_controle['Infrações ']
        for row in ws_infr.iter_rows(min_row=2, values_only=True):
            if not row or not row[9]:
                continue
            ait = str(row[9]).strip().upper()
            if not ait:
                continue
            placa = clean_placa(row[8])
            if not placa:
                continue
            
            data_infracao = format_timestamp_sql(row[10]) if len(row) > 10 else "NULL"
            desc_infracao = str(row[11]).strip() if len(row) > 11 and row[11] else "Infração de Trânsito"
            categoria = str(row[12]).strip().upper() if len(row) > 12 and row[12] else "MEDIA"
            cat_map = {
                "LEVE": "LEVE",
                "MÉDIA": "MEDIA",
                "MEDIA": "MEDIA",
                "GRAVE": "GRAVE",
                "GRAVÍSSIMA": "GRAVISSIMA",
                "GRAVISSIMA": "GRAVISSIMA"
            }
            gravidade = cat_map.get(categoria, "MEDIA")
            pontos = int(row[13]) if len(row) > 13 and str(row[13]).isdigit() else 0
            valor = clean_decimal(row[14], default=0.0) if len(row) > 14 else "0.00"
            matr = str(row[3]).strip() if len(row) > 3 and row[3] else None

            tec_subquery = f"(SELECT id FROM tecnicos WHERE matricula = {escape_sql(matr)} LIMIT 1)" if matr else "NULL"

            sql_statements.append(
                f"INSERT INTO multas_notificacoes (numero_ait, veiculo_id, tecnico_id, data_hora_infracao, descricao_infracao, gravidade, pontuacao, valor_original, status_pagamento) "
                f"SELECT {escape_sql(ait)}, v.id, {tec_subquery}, COALESCE({data_infracao}, CURRENT_TIMESTAMP), {escape_sql(desc_infracao)}, '{gravidade}', {pontos}, {valor}, 'PAGO' "
                f"FROM veiculos v WHERE v.placa = '{placa}' "
                f"ON CONFLICT (numero_ait) DO NOTHING;"
            )
    sql_statements.append("")

    # 9. Notificações Recentes de Multas (Multas notificação.csv)
    sql_statements.append("-- 9. Notificações Recentes com Prazos de Indicação de Condutor (Prevenção NIC)")
    csv_multas = os.path.join(DOCS_PLANILHAS, "Multas notificação.csv")
    if os.path.exists(csv_multas):
        with open(csv_multas, mode='r', encoding='latin1') as f:
            reader = csv.reader(f, delimiter=';')
            for i, row in enumerate(reader):
                if i == 0 or len(row) < 10:
                    continue
                placa = clean_placa(row[0])
                fez_indicacao = True if str(row[1]).strip().lower() == "sim" else False
                data_infr = row[2].strip()
                hora_infr = row[3].strip()
                dt_str = f"{data_infr} {hora_infr}" if data_infr else None
                data_hora_sql = format_timestamp_sql(dt_str)
                limite_hub = format_date_sql(row[4])
                limite_orgao = format_date_sql(row[5])
                ait = str(row[9]).strip().upper()
                link = str(row[10]).strip() if len(row) > 10 else None

                if ait and placa:
                    ind_bool = "TRUE" if fez_indicacao else "FALSE"
                    sql_statements.append(
                        f"INSERT INTO multas_notificacoes (numero_ait, veiculo_id, data_hora_infracao, data_limite_indicacao, condutor_indicado, link_documento, status_pagamento) "
                        f"SELECT {escape_sql(ait)}, v.id, COALESCE({data_hora_sql}, CURRENT_TIMESTAMP), {limite_orgao}, {ind_bool}, {escape_sql(link)}, 'PENDENTE' "
                        f"FROM veiculos v WHERE v.placa = '{placa}' "
                        f"ON CONFLICT (numero_ait) DO UPDATE SET "
                        f"data_limite_indicacao = EXCLUDED.data_limite_indicacao, "
                        f"condutor_indicado = EXCLUDED.condutor_indicado, "
                        f"link_documento = EXCLUDED.link_documento;"
                    )
    sql_statements.append("")

    # Escrever arquivo de migração V7
    with open(OUTPUT_SQL, mode="w", encoding="utf-8") as f:
        f.write("\n".join(sql_statements))

    print(f"Migração V7 gerada com sucesso em: {OUTPUT_SQL}")
    print(f"Total de comandos SQL: {len(sql_statements)}")

if __name__ == "__main__":
    run_etl()
